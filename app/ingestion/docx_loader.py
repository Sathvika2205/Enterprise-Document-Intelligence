from docx import Document

from app.ingestion.models import DocumentContent


def load_docx(file_path: str) -> list[DocumentContent]:
    document = Document(file_path)

    sections = []

    paragraph_text = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraph_text.append(text)

    if paragraph_text:
        sections.append(
            DocumentContent(
                text="\n".join(paragraph_text),
                source=file_path,
                file_type="docx",
                section="paragraphs",
            )
        )

    for table_index, table in enumerate(document.tables, start=1):
        rows = []

        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))

        if rows:
            sections.append(
                DocumentContent(
                    text="\n".join(rows),
                    source=file_path,
                    file_type="docx",
                    section=f"table_{table_index}",
                )
            )

    return sections